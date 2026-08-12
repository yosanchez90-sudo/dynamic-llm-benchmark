"""Generate docs/RESEARCH_QUESTIONS.md, a diffable/greppable markdown mirror
of docs/Research_questions.xlsx.

The xlsx is the source of truth for formatting (fonts, fills, borders,
column widths, freeze panes) and is edited by hand or via openpyxl; this
script only reads it and writes a plain-markdown copy of its cell values
so the workbook's content shows up in `git diff` and `grep`. Safe to
re-run any time after the xlsx changes; do not hand-edit the generated
.md file.

Usage:
    python scripts/build_research_questions_md.py
"""
from __future__ import annotations

import subprocess
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
XLSX_PATH = ROOT / "docs" / "Research_questions.xlsx"
MD_PATH = ROOT / "docs" / "RESEARCH_QUESTIONS.md"


def _git_short_sha() -> str:
    try:
        return subprocess.check_output(
            ["git", "rev-parse", "--short", "HEAD"], cwd=ROOT
        ).decode().strip()
    except Exception:
        return "unknown"


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    return str(value).replace("\n", "<br>").replace("|", "\\|")


def main() -> None:
    wb = openpyxl.load_workbook(XLSX_PATH, data_only=True)
    commit = _git_short_sha()

    lines = [
        "# Research questions workbook — markdown mirror",
        "",
        "Generated from `docs/Research_questions.xlsx` by "
        "`scripts/build_research_questions_md.py`.",
        f"Git commit at generation time: {commit}",
        "",
        "The xlsx is the source of truth for formatting; this file exists so "
        "the",
        "workbook's content is diffable and greppable in git. Do not "
        "hand-edit —",
        "edit the xlsx and regenerate.",
        "",
    ]

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        lines.append(f"## {sheet_name}")
        lines.append("")
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
        header = rows[0]
        ncols = len(header)
        lines.append("| " + " | ".join(_cell_text(h) for h in header) + " |")
        lines.append("|" + "|".join(["---"] * ncols) + "|")
        for row in rows[1:]:
            if all(v is None for v in row):
                continue
            lines.append("| " + " | ".join(_cell_text(v) for v in row) + " |")
        lines.append("")

    MD_PATH.write_text("\n".join(lines) + "\n")
    print(f"Wrote {MD_PATH.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
