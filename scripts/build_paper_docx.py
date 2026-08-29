"""Build paper/paper.docx: a Word rendering of the compiled paper.tex,
matching its two-column IEEE layout, section/table/figure numbering, and
resolved citations -- for review and markup only.

CLAUDE.md hard rule: the LaTeX in paper/ is the single source of truth
for the paper; this .docx is a rendering of it, never an editing base.
If changes are made in the .docx, they need to be relayed back and
applied to paper/sections/*.tex directly, not treated as authoritative.

Built from paper.tex's own source, not a PDF-to-DOCX conversion: a raw
PDF-to-docx pass (tested with pdf2docx) merges words together with no
space ("Dynamic,Template-BasedBenchmarking...") wherever the PDF's own
kerning has no explicit inter-word gap glyph, which is common in
justified two-column IEEE text -- useless for actual editing. Reading
the .tex source directly guarantees correct spacing, at the cost of a
best-effort (not pixel-perfect) LaTeX-to-Word conversion: inline math
renders as approximated Unicode/plain text, not native Word equations.

Section, table, and figure numbers are read directly from paper.aux's
own \\newlabel/\\bibcite entries (the same numbers LaTeX itself resolved
for the compiled PDF), not re-derived here, so they can never drift from
what the PDF actually shows -- requires a fresh `make` in paper/ first.

This script OVERWRITES paper/paper.docx unconditionally by default. That
silently destroyed real hand-edits once already (confirmed 2026-08-29:
this script's own Table II fix was run without checking first, wiping a
prose revision the author had made in Word since the last sync -- the
whole revision had to be recovered from a docx the author re-sent). It
now refuses to overwrite a docx that was edited in a real Word session
since it was last written by this script (see _docx_was_hand_edited());
pass --force only once those edits are confirmed already relayed into
paper/sections/*.tex, e.g. right after this script itself just wrote the
file being force-overwritten.

Usage: python scripts/build_paper_docx.py [--force]
"""
from __future__ import annotations

import re
import sys
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Pt, Inches, RGBColor, Cm

ROOT = Path(__file__).resolve().parent.parent
PAPER = ROOT / "paper"
SECTIONS = PAPER / "sections"
FIGURES = ROOT / "analysis" / "figures"
OUT_PATH = PAPER / "paper.docx"

NAVY = RGBColor(0x1F, 0x3A, 0x5F)
DARK = RGBColor(0x1A, 0x1A, 0x1A)

# ---------------------------------------------------------------- loaders

def _load_macros() -> dict[str, str]:
    text = (PAPER / "values.tex").read_text()
    return dict(re.findall(r"\\newcommand\{\\(\w+)\}\{(.*?)\}\n", text))


def _load_bibcite() -> dict[str, str]:
    text = (PAPER / "paper.aux").read_text()
    return dict(re.findall(r"\\bibcite\{(\w+)\}\{(\d+)\}", text))


def _load_reflabels() -> dict[str, str]:
    text = (PAPER / "paper.aux").read_text()
    return dict(re.findall(r"\\newlabel\{([^}]+)\}\{\{([^{}]+)\}", text))


def _abbreviate_name(full_name: str) -> str:
    """"Oier Lopez de Lacalle" -> "O. L. de Lacalle": given names become
    initials, the family name (BibTeX's "von Last" convention -- a
    lowercase-initial token and everything after it, else just the final
    token) is kept whole. A hyphenated given name gets one initial per
    part ("See-Kiong" -> "S.-K."), matching IEEEtran.bst's own author
    formatting exactly, which this docx must match (2026-08-29: the
    author found the docx's un-abbreviated "Oscar Sainz, Jon Ander
    Campos, ..." wasted space compared to the compiled PDF's "O. Sainz,
    J. A. Campos, ..." and wants both the same -- the PDF's shorter
    style, not a hand-shortened docx-only list that could drift from
    references.bib)."""
    tokens = full_name.split()
    von_start = next((i for i, t in enumerate(tokens) if i > 0 and t[0].islower()), None)
    given, last = (tokens[:-1], tokens[-1]) if von_start is None else (tokens[:von_start], " ".join(tokens[von_start:]))
    def _initial(tok: str) -> str:
        return "-".join(f"{part[0]}." for part in tok.split("-") if part)
    return f"{' '.join(_initial(g) for g in given)} {last}".strip()


def _format_authors(raw: str) -> str:
    """BibTeX's " and "-separated author list -> "A, B, and C" (or "A and
    B" for two, or "A et al." if the list ends in the "others" keyword),
    with every given name abbreviated to initials to match IEEEtran.bst's
    own reference formatting in the compiled PDF."""
    if not raw:
        return raw
    names = [n.strip() for n in raw.split(" and ")]
    et_al = names and names[-1].lower() == "others"
    if et_al:
        names = names[:-1]
    abbrev = [_abbreviate_name(n) for n in names]
    if len(abbrev) == 1:
        joined = abbrev[0]
    elif len(abbrev) == 2:
        joined = f"{abbrev[0]} and {abbrev[1]}"
    else:
        joined = ", ".join(abbrev[:-1]) + f", and {abbrev[-1]}"
    return joined + (" et al." if et_al else "")


def _bib_field(block: str, field: str) -> str:
    m = re.search(rf"\b{field}\s*=\s*\{{\{{?(.*?)\}}?\}},?\s*\n", block, re.DOTALL)
    return m.group(1).strip() if m else ""


def _load_references(bibcite: dict[str, str]) -> list[str]:
    """Return reference strings indexed 0..N-1 for numbers 1..N, formatted
    plainly (author, "title," venue, year, URL) -- not IEEE-exact, but
    complete and readable. Parses each @type{key, ...} block's fields
    independently (not by fixed field order): this project's own
    references.bib mixes booktitle/journal/howpublished and sometimes
    inserts eprint/archivePrefix between year and url."""
    bib_text = (PAPER / "references.bib").read_text()
    blocks = re.split(r"\n(?=@)", bib_text)
    by_key: dict[str, str] = {}
    for block in blocks:
        km = re.match(r"@\w+\{(\w+),", block)
        if not km:
            continue
        key = km.group(1)
        author = _format_authors(_bib_field(block, "author"))
        title = _bib_field(block, "title")
        venue = _bib_field(block, "booktitle") or _bib_field(block, "journal") or _bib_field(block, "howpublished")
        year = _bib_field(block, "year")
        url = _bib_field(block, "url")
        author_sep = "" if author.endswith(".") else "."  # avoid "et al.." (author already ends in a period)
        by_key[key] = f"{author}{author_sep} \u201c{title}.\u201d {venue}, {year}. {url}"

    ordered: list[tuple[int, str]] = [
        (int(num), by_key[key]) for key, num in bibcite.items() if key in by_key
    ]
    ordered.sort()
    return [r for _, r in ordered]


# ---------------------------------------------------------- LaTeX -> runs

_MATH_SUBS = [
    (r"\\tau", "\u03c4"), (r"\\rho", "\u03c1"), (r"\\geq", "\u2265"),
    (r"\\leq", "\u2264"), (r"\\times", "\u00d7"), (r"\\leftrightarrow", "\u2194"),
    (r"\\rightarrow", "\u2192"), (r"\\pm", "\u00b1"), (r"\\emph\{([^{}]*)\}", r"\1"),
    (r"\\text\{([^{}]*)\}", r"\1"), (r"\\dagger", "\u2020"),
]


def _strip_math(s: str) -> str:
    """Best-effort: turn $...$ math into readable plain text/Unicode, not a
    native Word equation. Handles the patterns actually used in this
    paper (Greek letters, comparisons, simple sub/superscripts)."""
    def repl(m: re.Match) -> str:
        inner = m.group(1)
        for pat, sub in _MATH_SUBS:
            inner = re.sub(pat, sub, inner)
        inner = re.sub(r"_\{([^{}]*)\}", r"_\1", inner)   # n_{def} -> n_def
        inner = re.sub(r"\^\{([^{}]*)\}", r"^\1", inner)  # x^{-4} -> x^-4
        inner = inner.replace("{", "").replace("}", "")
        return inner
    return re.sub(r"\$([^$]*)\$", repl, s)


def _normalize_escapes(text: str) -> str:
    """The handful of LaTeX escapes this project's own prose uses, outside
    of \\cite/\\ref/\\label/\\textbf/\\texttt handling (those need
    surrounding context, so they stay in _add_latex_runs)."""
    text = text.replace("\\%", "\ufe5f").replace("\\&", "\ufe60").replace("\\_", "\ufe63")
    text = text.replace("\\{", "{").replace("\\}", "}")
    text = text.replace("---", "\u2014").replace("--", "\u2013")
    text = text.replace("\\S", "\u00a7").replace("~", " ")
    text = text.replace("\\ ", " ")  # LaTeX controlled space (e.g. "vs.\ ", "Exec.\ =")
    text = re.sub(r"\^(?!\{)", "", text)  # stray bare ^ left by $^\dagger$-style superscripts
    return text


def _clean_title(text: str, macros: dict[str, str]) -> str:
    """Plain-text cleanup for heading/title strings that don't need
    _add_latex_runs's run-tokenization (no \\textbf/\\texttt inside a
    \\subsection{} title in this paper, but macros and escapes do appear,
    e.g. \\subsection{H1 Test 1: DyMuT-L4 vs.\\ ULT ...})."""
    for name, val in macros.items():
        text = text.replace(f"\\{name}{{}}", val)
    text = _strip_math(text)
    text = _normalize_escapes(text)
    return _restore_escapes(text)


def _add_latex_runs(p, text: str, macros: dict[str, str], bibcite: dict[str, str],
                     reflabels: dict[str, str]) -> None:
    """Parse a single paragraph's worth of paper LaTeX into styled runs.
    Handles \\textbf, \\emph, \\texttt, \\cite, \\ref/\\S, macros, $..$
    math, and the handful of escapes this project's own prose uses
    (\\%, \\&, \\_, --). Not a general LaTeX parser -- only what this
    paper's own sections/*.tex actually contain."""
    # Macro substitution first, so nothing downstream needs to know macros exist.
    for name, val in macros.items():
        text = text.replace(f"\\{name}{{}}", val)
    text = _strip_math(text)
    text = _normalize_escapes(text)
    text = re.sub(r"% cite:\w+.*$", "", text)  # drop trailing citation-gate markers

    def cite_repl(m: re.Match) -> str:
        keys = [k.strip() for k in m.group(1).split(",")]
        nums = [bibcite.get(k, "?") for k in keys]
        return "[" + ", ".join(nums) + "]"

    text = re.sub(r"\\cite\{([^{}]+)\}", cite_repl, text)
    text = re.sub(r"\\ref\{([^{}]+)\}", lambda m: reflabels.get(m.group(1), "?"), text)
    text = re.sub(r"\\label\{[^{}]+\}", "", text)
    text = re.sub(r"\\emph\{([^{}]*)\}", r"\1", text)

    # Tokenize on \textbf{...} / \texttt{...} so those spans get their own
    # run with formatting; everything else is a plain run.
    token_re = re.compile(r"\\textbf\{([^{}]*)\}|\\texttt\{([^{}]*)\}")
    pos = 0
    for m in token_re.finditer(text):
        if m.start() > pos:
            _plain_run(p, text[pos:m.start()])
        if m.group(1) is not None:
            r = p.add_run(_restore_escapes(m.group(1)))
            r.bold = True
        else:
            r = p.add_run(_restore_escapes(m.group(2)))
            r.font.name = "Consolas"
        pos = m.end()
    if pos < len(text):
        _plain_run(p, text[pos:])


def _restore_escapes(text: str) -> str:
    return text.replace("\ufe5f", "%").replace("\ufe60", "&").replace("\ufe63", "_")


def _plain_run(p, text: str) -> None:
    text = _restore_escapes(text)
    if text:
        p.add_run(text)


# --------------------------------------------------------------- tables

def _strip_line_comments(tex: str) -> str:
    """Drop a trailing, unescaped LaTeX '%...' comment from every line.
    Without this, a row-terminating '\\\\ % cite:external ...' marker
    (this project's citation-gate suppression convention) gets absorbed
    into the *next* row's first cell by _parse_tabular_rows's split on
    '\\\\', corrupting that cell's rendered text -- confirmed in Table IV
    (H2-supplementary), where three Model cells rendered the comment text
    instead of "Qwen 1.5B"/"Qwen 7B". '\\%' (an escaped, literal percent
    sign) must not be treated as a comment start."""
    return "\n".join(re.sub(r"(?<!\\)%.*$", "", line) for line in tex.splitlines())


def _flatten_shortstack(tex: str) -> str:
    """Replace \\shortstack[pos]{a\\\\b} with its content, '\\\\' -> ' '.
    Real LaTeX scopes the '\\\\' inside \\shortstack's braces correctly (it
    never reaches the tabular row-separator parser), but this script's own
    _parse_tabular_rows/_header_row split naively on every '\\\\' with no
    brace-awareness -- an un-flattened \\shortstack header cell (e.g. a
    two-line "Cov.-\\\\adj." column head) gets misread as ending the row
    early, corrupting every cell after it. Confirmed: Table II's header
    row collapsed from 11 cells to 2 without this.

    A regex with a fixed nesting depth isn't enough here: some cells
    (e.g. \\shortstack{$n_{\\text{def}}$}) nest braces two levels deep
    ({\\text{def}} inside the $..$), so this scans for the matching close
    brace by tracking depth instead of assuming a brace count."""
    out = []
    i = 0
    pattern = re.compile(r"\\shortstack(?:\[[a-z]\])?\{")
    while True:
        m = pattern.search(tex, i)
        if not m:
            out.append(tex[i:])
            break
        out.append(tex[i:m.start()])
        depth = 1
        j = m.end()
        while j < len(tex) and depth > 0:
            if tex[j] == "{":
                depth += 1
            elif tex[j] == "}":
                depth -= 1
            j += 1
        content = tex[m.end():j - 1]
        out.append(content.replace("\\\\", " "))
        i = j
    return "".join(out)


def _parse_tabular_rows(tex: str) -> list[list[str]]:
    """Extract data rows (as raw cell strings) from the first
    tabular/tabularx environment in tex, between \\midrule and
    \\bottomrule, skipping \\toprule/\\midrule/rule lines and a trailing
    \\multicolumn note row."""
    body = re.search(r"\\midrule(.*?)\\bottomrule", tex, re.DOTALL)
    if not body:
        return []
    rows = []
    for line in body.group(1).split("\\\\"):
        line = line.strip()
        if not line or line.startswith("\\midrule") or line.startswith("\\cmidrule"):
            continue
        cells = [c.strip() for c in line.split("&")]
        rows.append(cells)
    return rows


def _header_row(tex: str) -> list[str]:
    m = re.search(r"\\toprule(.*?)\\midrule", tex, re.DOTALL)
    if not m:
        return []
    span = m.group(1)
    # \cmidrule(lr){2-6} sub-header rules carry no trailing \\ of their
    # own (they sit between the multicolumn group-header row and the
    # actual column-name row), so splitting on \\ alone merges them into
    # whichever row follows; drop them by line first.
    span = "\n".join(l for l in span.splitlines() if not l.strip().startswith("\\cmidrule"))
    # last non-empty *data* row before \midrule is the column-name row;
    # a multicolumn group-header row (e.g. spanning "DyMuT (pooled)" /
    # "ULT (core_400)") is dropped too -- editing draft, not pixel-exact.
    lines = [l.strip() for l in span.split("\\\\") if l.strip()]
    lines = [l for l in lines if "\\multicolumn" not in l]
    if not lines:
        return []
    return [c.strip() for c in lines[-1].split("&")]


def _add_docx_table(doc, tex: str, macros, bibcite, reflabels, caption: str) -> None:
    tex = _strip_line_comments(tex)
    tex = _flatten_shortstack(tex)
    header = _header_row(tex)
    rows = _parse_tabular_rows(tex)
    if not header or not rows:
        return
    n_cols = len(header)
    table = doc.add_table(rows=1, cols=n_cols)
    table.style = "Light Grid Accent 1" if "Light Grid Accent 1" in [s.name for s in doc.styles] else "Table Grid"
    for c, cell_tex in enumerate(header):
        p = table.rows[0].cells[c].paragraphs[0]
        _add_latex_runs(p, cell_tex, macros, bibcite, reflabels)
        for r in p.runs:
            r.bold = True
            r.font.size = Pt(9)
    for row_vals in rows:
        cells = table.add_row().cells
        for c, cell_tex in enumerate(row_vals[:n_cols]):
            p = cells[c].paragraphs[0]
            _add_latex_runs(p, cell_tex, macros, bibcite, reflabels)
            for r in p.runs:
                r.font.size = Pt(9)
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _add_latex_runs(cap, caption, macros, bibcite, reflabels)
    for r in cap.runs:
        r.italic = True
        r.font.size = Pt(9.5)
    doc.add_paragraph()


# ------------------------------------------------------------- sections

def _set_two_columns(section) -> None:
    """python-docx's own section template already ships a default
    <w:cols w:space="720"/> element; appending a second <w:cols> (as
    opposed to replacing the existing one) leaves two competing elements
    in sectPr and Word honours the first (single-column) one, silently
    discarding the two-column setting -- confirmed by inspecting the
    saved XML directly."""
    sect_pr = section._sectPr
    existing = sect_pr.find(qn("w:cols"))
    if existing is not None:
        sect_pr.remove(existing)
    cols = OxmlElement("w:cols")
    cols.set(qn("w:num"), "2")
    cols.set(qn("w:space"), "425")  # ~0.3in gutter, IEEE-ish
    sect_pr.append(cols)


def _numbered_heading(doc, number: str, title: str, level: int) -> None:
    style_name = f"Heading {min(level, 3)}"
    h = doc.add_heading(level=level)
    r = h.add_run(f"{number}  {title}")
    r.font.size = Pt(13 if level == 1 else 11.5)
    r.font.bold = True
    r.font.color.rgb = NAVY
    r.font.name = "Calibri"


def _section_title(path: Path) -> str:
    """Read a section's own \\section{...} text directly, instead of a
    hardcoded fallback string duplicated here -- a hardcoded title can
    drift silently from the real one and never gets caught (confirmed
    2026-08-29: this file's own body_sections list still said "Related
    Work" after the actual section had been renamed "Literature review"
    in 02_related_work.tex, so the docx kept showing the old title)."""
    m = re.search(r"\\section\{(.*?)\}", path.read_text())
    return m.group(1) if m else "?"


def _walk_section_file(doc, path: Path, macros, bibcite, reflabels) -> None:
    text = path.read_text()
    lines = text.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        if not line or line.startswith("%") or line.startswith("\\section") or line.startswith("\\label"):
            i += 1
            continue
        if line.startswith("\\subsection{"):
            m = re.search(r"\\subsection\{(.*?)\}", line)
            title = m.group(1) if m else "?"
            # next non-empty line, if \label{...}, gives the resolved number
            j = i + 1
            num = "?"
            if j < len(lines) and lines[j].strip().startswith("\\label{"):
                lm = re.search(r"\\label\{([^{}]+)\}", lines[j])
                if lm:
                    num = reflabels.get(lm.group(1), "?")
                i = j
            _numbered_heading(doc, num, _clean_title(title, macros), level=2)
            i += 1
            continue
        if line.startswith("\\begin{figure"):
            j = i
            block = []
            while j < len(lines) and "\\end{figure" not in lines[j]:
                block.append(lines[j]); j += 1
            block_text = "\n".join(block)
            img_m = re.search(r"\\includegraphics(?:\[[^\]]*\])?\{([^{}]+)\}", block_text)
            # MULTILINE only, no DOTALL: this project writes each \caption{}
            # as a single physical source line (CLAUDE.md convention), so a
            # greedy same-line match is safe and, critically, does NOT run
            # past the caption's own closing brace onto the following
            # \label{fig:...} line -- confirmed this was happening with a
            # DOTALL match (greedy `.*` swallowed through to \label{fig:
            # onset}'s own closing brace, leaving an unmatched, unstripped
            # "\label{fig:onset" fragment in the rendered caption text).
            cap_m = re.search(r"\\caption\{(.*)\}\s*$", block_text, re.MULTILINE)
            if img_m:
                png = FIGURES / f"{img_m.group(1)}.png"
                if png.exists():
                    doc.add_picture(str(png), width=Inches(3.4))
                    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            if cap_m:
                cap_text = cap_m.group(1)
                # This section now has three figures (fig:onset, fig:h3tiereffect,
                # fig:h5stability), not one -- read each figure's own \label{}
                # from its block instead of the fig:onset hardcode that was
                # correct back when this was the only figure in the paper.
                label_m = re.search(r"\\label\{(fig:[^{}]+)\}", block_text)
                fig_num = reflabels.get(label_m.group(1), "?") if label_m else "?"
                cap = doc.add_paragraph()
                cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r = cap.add_run(f"Fig. {fig_num}. ")
                r.italic = True
                r.bold = True
                r.font.size = Pt(9)
                p2 = cap
                _add_latex_runs(p2, cap_text, macros, bibcite, reflabels)
                for run in p2.runs[1:]:
                    run.italic = True
                    run.font.size = Pt(9)
            i = j + 1
            continue
        if line.startswith("\\input{table1}") or line.startswith("\\input{table_results}"):
            tex_name = "table1.tex" if "table1" in line else "table_results.tex"
            tex = (PAPER / tex_name).read_text()
            cap_m = re.search(r"\\caption\{(.*?)\}\s*\\label\{(tab:\w+)\}", tex, re.DOTALL)
            if cap_m:
                cap_text = cap_m.group(1)
                label = cap_m.group(2)
                num = reflabels.get(label, "?")
                doc.add_paragraph()
                title_p = doc.add_paragraph()
                title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r = title_p.add_run(f"TABLE {num}")
                r.bold = True
                r.font.size = Pt(10)
                if tex_name == "table_results.tex":
                    # The PDF's multicolumn group header ("DyMuT (pooled)"
                    # spanning cols 2-6, "ULT (core_400)" spanning 7-11)
                    # doesn't survive this simplified single-header-row
                    # parse; state it in a note instead of losing it silently.
                    note_p = doc.add_paragraph()
                    note_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    nr = note_p.add_run("Columns 2–6: DyMuT (pooled). Columns 7–11: ULT (core_400).")
                    nr.italic = True
                    nr.font.size = Pt(8.5)
                _add_docx_table(doc, tex, macros, bibcite, reflabels, cap_text)
            i += 1
            continue
        if line.startswith("\\begin{table"):
            j = i
            block = []
            while j < len(lines) and "\\end{table" not in lines[j]:
                block.append(lines[j]); j += 1
            block_text = "\n".join(block)
            cap_m = re.search(r"\\caption\{(.*?)\}\s*\\label\{(tab:\w+)\}", block_text, re.DOTALL)
            if cap_m:
                cap_text = cap_m.group(1)
                label = cap_m.group(2)
                num = reflabels.get(label, "?")
                doc.add_paragraph()
                title_p = doc.add_paragraph()
                title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r = title_p.add_run(f"TABLE {num}")
                r.bold = True
                r.font.size = Pt(10)
                _add_docx_table(doc, block_text, macros, bibcite, reflabels, cap_text)
            i = j + 1
            continue
        # plain paragraph (this project writes one paragraph per physical line)
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        _add_latex_runs(p, line, macros, bibcite, reflabels)
        i += 1


# ------------------------------------------------------------------ main

def build() -> None:
    macros = _load_macros()
    bibcite = _load_bibcite()
    reflabels = _load_reflabels()
    references = _load_references(bibcite)

    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(4)

    section = doc.sections[0]
    section.page_width = Cm(21.59)   # US letter, matching \documentclass[letterpaper]
    section.page_height = Cm(27.94)
    section.left_margin = Cm(1.9)
    section.right_margin = Cm(1.9)
    section.top_margin = Cm(1.9)
    section.bottom_margin = Cm(1.9)

    # ---- title block (single column) ----
    title_tex = (
        "Dynamic, Template-Based Benchmarking Resists "
        "Contamination-Induced Score Inflation in LLM Unit-Test Generation"
    )
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(title_tex)
    r.bold = True
    r.font.size = Pt(20)
    r.font.color.rgb = NAVY

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Yohan Sneider Sanchez Vega")
    r.font.size = Pt(12)
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.add_run(
        "School of Electronics, Electrical Engineering and Computer Science\n"
        "Queen's University Belfast, Belfast, United Kingdom\n"
        "ysanchezvega01@qub.ac.uk"
    )
    r2.font.size = Pt(10)

    abstract_tex = (SECTIONS / "00_abstract.tex").read_text().strip()
    doc.add_paragraph()
    p = doc.add_paragraph()
    r = p.add_run("Abstract\u2014")
    r.bold = True
    r.italic = True
    _add_latex_runs(p, abstract_tex, macros, bibcite, reflabels)
    for run in p.runs[1:]:
        run.italic = True
        run.font.size = Pt(9.5)

    p = doc.add_paragraph()
    r = p.add_run(
        "Index Terms\u2014data contamination, dynamic benchmarking, "
        "large language models, mutation testing, unit test generation"
    )
    r.italic = True
    r.font.size = Pt(9.5)

    # ---- switch to two-column for the body ----
    new_sect = doc.add_section(WD_SECTION.CONTINUOUS)
    new_sect.left_margin = section.left_margin
    new_sect.right_margin = section.right_margin
    new_sect.top_margin = section.top_margin
    new_sect.bottom_margin = section.bottom_margin
    _set_two_columns(new_sect)

    body_sections = [
        (SECTIONS / "01_introduction.tex", "sec:introduction"),
        (SECTIONS / "02_related_work.tex", "sec:related-work"),
        (SECTIONS / "03_methods.tex", "sec:methods"),
        (SECTIONS / "04_results.tex", "sec:results"),
        (SECTIONS / "05_discussion.tex", "sec:discussion"),
        (SECTIONS / "05a_ethics.tex", "sec:ethics"),
        (SECTIONS / "06_conclusion.tex", "sec:conclusion"),
    ]
    for path, label in body_sections:
        num = reflabels.get(label, "?")
        _numbered_heading(doc, num, _section_title(path), level=1)
        _walk_section_file(doc, path, macros, bibcite, reflabels)

    # No Acknowledgment section: paper.tex itself no longer \inputs
    # sections/07_acknowledgment.tex (cut during the page-budget pass to
    # hit the 8-page limit -- the file is still on disk but unreferenced).
    # Rendering it here anyway would show content in the docx that isn't
    # in the actual compiled paper, which defeats this script's whole
    # point (matching the PDF exactly) and would give an edit-diff tool
    # phantom content with no real source line to route changes back to.

    # ---- references ----
    ref_h = doc.add_heading(level=1)
    r = ref_h.add_run("References")
    r.font.size = Pt(13); r.font.bold = True; r.font.color.rgb = NAVY
    for n, ref in enumerate(references, start=1):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.left_indent = Cm(0.5)
        p.paragraph_format.first_line_indent = Cm(-0.5)
        run = p.add_run(f"[{n}] {ref}")
        run.font.size = Pt(8.5)

    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    if _docx_was_hand_edited(OUT_PATH) and "--force" not in sys.argv:
        raise SystemExit(
            f"Refusing to overwrite {OUT_PATH.relative_to(ROOT)}: its "
            f"docProps/core.xml shows a real edit (lastModifiedBy set, or "
            f"a revision/edit-time count beyond a fresh save), meaning "
            f"someone edited it in Word since it was last written by this "
            f"script. Regenerating now would silently discard that work "
            f"(this happened for real, 2026-08-29). Relay those edits "
            f"into paper/sections/*.tex first (read the docx, diff it "
            f"against a build into a scratch path, apply the real "
            f"differences), then re-run with --force."
        )
    doc.save(OUT_PATH)
    print(f"Saved: {OUT_PATH.relative_to(ROOT)} ({len(references)} references)")


def _docx_was_hand_edited(path: Path) -> bool:
    """True if path exists and its docProps/core.xml shows a real Word
    editing session (lastModifiedBy set, or revision > 1) rather than
    being exactly what this script last wrote. python-docx's own
    Document() template leaves lastModifiedBy empty and revision at 1;
    Word sets both the moment a human opens and saves the file."""
    if not path.exists():
        return False
    try:
        with zipfile.ZipFile(path) as z:
            core = z.read("docProps/core.xml").decode("utf-8", errors="replace")
    except (KeyError, zipfile.BadZipFile, FileNotFoundError, OSError):
        return False
    modified_by = re.search(r"<cp:lastModifiedBy>(.*?)</cp:lastModifiedBy>", core)
    revision = re.search(r"<cp:revision>(\d+)</cp:revision>", core)
    return bool((modified_by and modified_by.group(1).strip())
                or (revision and int(revision.group(1)) > 1))


if __name__ == "__main__":
    build()
